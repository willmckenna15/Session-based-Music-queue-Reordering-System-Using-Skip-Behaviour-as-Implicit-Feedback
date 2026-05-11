import pandas as pd
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Border, Side, Font, PatternFill
from openpyxl.chart.legend import Legend
from openpyxl.chart.layout import Layout, ManualLayout

output_path = "../Volunteer Data/Adhoc Analyses"

while True:
    user_id = input("Volunteer's Name: ")
    try:
        df = pd.read_csv(f"../Volunteer Data/RAW/Streaming_History_{user_id}.csv").copy()
    except FileNotFoundError:
        print(f"No file found for name: {user_id}")
        continue
    break

df = df.rename(columns={
    "master_metadata_track_name": "Track Name",
    "master_metadata_album_artist_name": "Artist Name"
})

df["ts"] = pd.to_datetime(df["ts"])
df["date"] = df["ts"].dt.date
df["time"] = df["ts"].dt.time

def Artist_Plays():
    Artist = input("What artist do you want to investigate? ")
    df["Month"] = df["ts"].dt.to_period("M").astype(str)
    artist_listens = df[df["Artist Name"] == Artist].groupby("Month").size().rename("Amount of listens").reset_index()
    all_months = pd.period_range(start=artist_listens["Month"].min(), end=artist_listens["Month"].max(), freq="M").astype(str)
    artist_listens = artist_listens.set_index("Month").reindex(all_months, fill_value=0).reset_index()
    artist_listens.columns = ["Month", "Amount of listens"]
    num_rows = len(artist_listens) + 1  # +1 for header row
    

    with pd.ExcelWriter(f"{output_path}/{user_id}'s {Artist} plays.xlsx", engine="openpyxl") as writer:
        artist_listens.to_excel(writer, sheet_name=f"{Artist} listens", index=False)

        worksheet = writer.sheets[f"{Artist} listens"]
        worksheet.sheet_state = "hidden"

        chart = LineChart()
        chart.title = f"Listens of {Artist} by month"
        chart.y_axis.title = "Listens"
        chart.x_axis.title = "Month"
        chart.width = 60
        chart.height = 20
        chart.legend = None
        chart.varyColors = False
        chart.x_axis.number_format = "yyyy-mm"
        chart.x_axis.tickLblPos = "low"
        chart.y_axis.numFmt = "0"
        chart.y_axis.tickLblPos = "nextTo"

        data = Reference(worksheet, min_col=2, max_col=2, min_row=1, max_row=num_rows)
        labels = Reference(worksheet, min_col=1, max_col=1, min_row=2, max_row=num_rows)

        chart.add_data(data, titles_from_data=True, from_rows=False)
        chart.set_categories(labels)
        chart.x_axis.tickLblSkip = 2
        chart.y_axis.scaling.min = 0
        chart.y_axis.majorGridlines = None
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        chart.style = 2
        chart.series[0].graphicalProperties.line.solidFill = "4472C4"
        chart.series[0].graphicalProperties.line.width = 25000

        chart_ws = writer.book.create_sheet("Listens by Month")
        chart_ws.add_chart(chart, "A1")

def main():
    functions = {
        "Artist Plays": Artist_Plays,
    }

    while True:
        print("Available functions:")
        for i, name in enumerate(functions, 1):
            print(f"  {i}. {name}")

        choice = input("\nWhat function would you like to run? ").strip()

        try:
            index = int(choice)
            selected = list(functions.values())[index - 1]
            selected()
        except (ValueError, IndexError):
            print(f"Invalid choice: '{choice}', please enter a number between 1 and {len(functions)}")
            continue
        break

if __name__ == "__main__":
    main()